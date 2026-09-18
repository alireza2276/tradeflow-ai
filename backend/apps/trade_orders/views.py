from rest_framework import status, viewsets
from rest_framework.views import APIView
from django.shortcuts import get_object_or_404
from django.core.exceptions import ValidationError as DjangoValidationError
from rest_framework.exceptions import ValidationError as DRFValidationError
from django.db import transaction

from rest_framework import status, viewsets
from rest_framework.response import Response
from django.db.models.deletion import ProtectedError


from apps.trade_orders.models import (
    CurrencyPurchase,
    PaymentInstrument,
    RegistrationOrder,
    ShipmentPart,
    RegulatoryRule,
    DeadlineExtension,
    CustomsClearance,
    RegulatoryDeadline,
)

from apps.trade_orders.serializers import (
    CurrencyPurchaseSerializer,
    PaymentInstrumentSerializer,
    RegistrationOrderSerializer,
    ShipmentPartSerializer,
    RegulatoryRuleSerializer,
    DeadlineExtensionSerializer,
    CustomsClearanceSerializer,
    RegulatoryDeadlineSerializer,
)

from apps.trade_orders.services.payment_instrument_service import (
    create_payment_instrument,
    update_payment_instrument,
)

from apps.trade_orders.services.dashboard_service import (
    get_dashboard_summary,
)

from apps.authentication.permissions import (
    CanViewTradeDashboard,
    TradeFlowModelPermissions,
)

from apps.workflows.services.submission_service import (
    submit_currency_purchase_correction,
    submit_currency_purchase_create,
    submit_currency_purchase_void,
    submit_shipment_part_correction,
    submit_shipment_part_create,
    submit_shipment_part_void,
)


from apps.trade_orders.services.compliance_service import (
    extend_deadline,
    refresh_obligation_status,
    mark_obligation_settled,
)
import csv
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from django.http import StreamingHttpResponse
from rest_framework.decorators import action

class RegistrationOrderViewSet(viewsets.ModelViewSet):

    permission_classes = (
        TradeFlowModelPermissions,
    )

    queryset = (
        RegistrationOrder.objects
        .select_related("company")
        .all()
    )

    serializer_class = RegistrationOrderSerializer

    search_fields = (
        "order_number",
        "company__name",
        "company__national_id",
        "currency",
    )

    filterset_fields = (
        "currency",
        "is_active",
    )

    ordering_fields = (
        "created_at",
        "registered_amount",
        "order_number",
    )

    @action(
        detail=False,
        methods=["get"],
        url_path="export",
    )
    def export_csv(self, request):
        queryset = self.filter_queryset(
            self.get_queryset()
        )

        class Echo:
            def write(self, value):
                return value

        pseudo_buffer = Echo()
        writer = csv.writer(pseudo_buffer)

        def safe_csv_value(value):
            if value is None:
                return ""

            text = str(value)

            if text.startswith(
                ("=", "+", "-", "@")
            ):
                return "'" + text

            return text

        def generate_rows():
            yield "\ufeff"

            yield writer.writerow([
                "Order Number",
                "Company",
                "National ID",
                "Registered Amount",
                "Currency",
                "Status",
            ])

            for order in queryset.iterator(
                chunk_size=1000
            ):
                yield writer.writerow([
                    safe_csv_value(
                        order.order_number
                    ),
                    safe_csv_value(
                        order.company.name
                    ),
                    safe_csv_value(
                        order.company.national_id
                    ),
                    safe_csv_value(
                        order.registered_amount
                    ),
                    safe_csv_value(
                        order.currency
                    ),
                    (
                        "Active"
                        if order.is_active
                        else "Inactive"
                    ),
                ])

        response = StreamingHttpResponse(
            generate_rows(),
            content_type=(
                "text/csv; charset=utf-8"
            ),
        )

        response[
            "Content-Disposition"
        ] = (
            'attachment; '
            'filename="registration-orders.csv"'
        )

        response[
            "X-Content-Type-Options"
        ] = "nosniff"

        return response

    def update(self, request, *args, **kwargs):
        partial = kwargs.pop("partial", False)

        with transaction.atomic():
            queryset = self.filter_queryset(
                self.get_queryset().select_for_update()
            )

            registration_order = get_object_or_404(
                queryset,
                pk=kwargs["pk"],
            )

            self.check_object_permissions(
                request,
                registration_order,
            )

            serializer = self.get_serializer(
                registration_order,
                data=request.data,
                partial=partial,
            )

            serializer.is_valid(
                raise_exception=True,
            )

            self.perform_update(
                serializer,
            )

            if getattr(
                    registration_order,
                    "_prefetched_objects_cache",
                    None,
            ):
                registration_order._prefetched_objects_cache = {}

            return Response(
                serializer.data
            )

    def destroy(self, request, *args, **kwargs):
        registration_order = self.get_object()

        try:
            registration_order.delete()
        except ProtectedError:
            return Response(
                {
                    "detail": (
                        "This registration order cannot be deleted "
                        "because it has related payment instruments "
                        "or currency purchases."
                    )
                },
                status=status.HTTP_409_CONFLICT,
            )

        return Response(
            status=status.HTTP_204_NO_CONTENT
        )
class PaymentInstrumentViewSet(viewsets.ModelViewSet):

    permission_classes = (
        TradeFlowModelPermissions,
    )

    queryset = (
        PaymentInstrument.objects
        .select_related(
            "registration_order",
            "registration_order__company",
        )
        .all()
    )

    serializer_class = PaymentInstrumentSerializer

    search_fields = (
        "instrument_number",
        "registration_order__order_number",
        "registration_order__company__name",
        "registration_order__company__national_id",
    )

    ordering_fields = (
        "created_at",
        "instrument_number",
    )

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(
            data=request.data,
        )
        serializer.is_valid(
            raise_exception=True,
        )

        registration_order = serializer.validated_data[
            "registration_order"
        ]

        instrument_number = serializer.validated_data[
            "instrument_number"
        ]

        payment_instrument = create_payment_instrument(
            registration_order=registration_order,
            instrument_number=instrument_number,
            operation_type=serializer.validated_data.get("operation_type", PaymentInstrument.OperationType.REMITTANCE),
            issue_date=serializer.validated_data.get("issue_date"),
        )

        output_serializer = self.get_serializer(
            payment_instrument,
        )

        return Response(
            output_serializer.data,
            status=status.HTTP_201_CREATED,
        )

    def update(self, request, *args, **kwargs):
        partial = kwargs.pop(
            "partial",
            False,
        )

        payment_instrument = self.get_object()

        serializer = self.get_serializer(
            payment_instrument,
            data=request.data,
            partial=partial,
        )

        serializer.is_valid(
            raise_exception=True,
        )

        data = serializer.validated_data

        registration_order = data.get(
            "registration_order",
            payment_instrument.registration_order,
        )

        if (
                registration_order.pk !=
                payment_instrument.registration_order_id
        ):
            raise DRFValidationError(
                {
                    "registration_order": (
                        "Registration order cannot be changed "
                        "after payment instrument creation."
                    )
                }
            )

        instrument_number = data.get(
            "instrument_number",
            payment_instrument.instrument_number,
        )

        try:
            updated_payment_instrument = (
                update_payment_instrument(
                    payment_instrument=payment_instrument,
                    instrument_number=instrument_number,
                    operation_type=data.get("operation_type", payment_instrument.operation_type),
                    issue_date=data.get("issue_date", payment_instrument.issue_date),
                )
            )
        except DjangoValidationError as exc:
            raise DRFValidationError(
                {
                    "detail": exc.messages,
                }
            )

        output_serializer = self.get_serializer(
            updated_payment_instrument,
        )

        return Response(
            output_serializer.data
        )

    @action(detail=True, methods=["post"], url_path="settle-obligation")
    def settle_obligation(self, request, *args, **kwargs):
        purchase = self.get_object()
        if not request.user.has_perm("trade_orders.settle_currencypurchase_obligation"):
            return Response(
                {"detail": "You do not have permission to settle FX obligations."},
                status=status.HTTP_403_FORBIDDEN,
            )
        try:
            updated = mark_obligation_settled(
                purchase=purchase, settled=True, actor=request.user,
                reason=request.data.get("reason", ""),
                reference=request.data.get("reference", ""),
            )
        except DjangoValidationError as exc:
            raise DRFValidationError({"detail": exc.messages})
        return Response(self.get_serializer(updated).data, status=status.HTTP_200_OK)

    @action(detail=True, methods=["post"], url_path="reopen-obligation")
    def reopen_obligation(self, request, *args, **kwargs):
        purchase = self.get_object()
        if not request.user.has_perm("trade_orders.settle_currencypurchase_obligation"):
            return Response(
                {"detail": "You do not have permission to reopen FX obligations."},
                status=status.HTTP_403_FORBIDDEN,
            )
        try:
            updated = mark_obligation_settled(
                purchase=purchase, settled=False, actor=request.user,
                reason=request.data.get("reason", ""),
                reference=request.data.get("reference", ""),
            )
        except DjangoValidationError as exc:
            raise DRFValidationError({"detail": exc.messages})
        return Response(self.get_serializer(updated).data, status=status.HTTP_200_OK)

    def destroy(self, request, *args, **kwargs):
        return Response(
            {
                "detail": (
                    "Payment instrument deletion is not allowed. "
                    "Use a correction or void workflow instead."
                )
            },
            status=status.HTTP_405_METHOD_NOT_ALLOWED,
        )

class CurrencyPurchaseViewSet(viewsets.ModelViewSet):
    permission_classes = (
        TradeFlowModelPermissions,
    )

    queryset = (
        CurrencyPurchase.objects
        .select_related(
            "registration_order",
            "registration_order__company",
            "registration_order__payment_instrument",
        )
        .all()
    )

    serializer_class = CurrencyPurchaseSerializer

    search_fields = (
        "registration_order__order_number",
        "registration_order__company__name",
        "registration_order__company__national_id",
        "registration_order__payment_instrument__instrument_number",
        "currency",
    )

    ordering_fields = (
        "purchase_date",
        "deadline",
        "amount",
        "created_at",
    )

    def get_queryset(self):
        queryset = super().get_queryset()
        params = self.request.query_params

        text_filters = (
            ("company", "registration_order__company__name__icontains"),
            ("national_id", "registration_order__company__national_id__icontains"),
            ("order_number", "registration_order__order_number__icontains"),
            ("instrument_number", "registration_order__payment_instrument__instrument_number__icontains"),
        )
        for param, lookup in text_filters:
            value = params.get(param, "").strip()
            if value:
                queryset = queryset.filter(**{lookup: value})

        currency = params.get("currency", "").strip().upper()
        if currency:
            queryset = queryset.filter(currency=currency)

        date_filters = (
            ("purchase_date_from", "purchase_date__gte"),
            ("purchase_date_to", "purchase_date__lte"),
            ("deadline_from", "deadline__gte"),
            ("deadline_to", "deadline__lte"),
        )
        for param, lookup in date_filters:
            value = params.get(param, "").strip()
            if value:
                queryset = queryset.filter(**{lookup: value})

        status_value = params.get("status", "").strip().lower()
        if status_value == "active":
            queryset = queryset.filter(is_void=False)
        elif status_value == "void":
            queryset = queryset.filter(is_void=True)

        for param, lookup in (
            ("amount_min", "amount__gte"),
            ("amount_max", "amount__lte"),
        ):
            value = params.get(param, "").strip()
            if value:
                try:
                    queryset = queryset.filter(
                        **{lookup: Decimal(value)}
                    )
                except InvalidOperation:
                    raise DRFValidationError(
                        {"detail": "Amount filters must be valid numbers."}
                    )
        return queryset

    @action(detail=False, methods=["get"], url_path="export-xlsx")
    def export_xlsx(self, request):
        queryset = self.filter_queryset(self.get_queryset())
        serializer = self.get_serializer(queryset, many=True)

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Currency Purchases"
        headers = [
            "Company", "National ID", "Registration Order",
            "Payment Instrument", "Purchase Sequence",
            "Registered Amount", "Purchase Amount", "Total Purchased",
            "Remaining To Purchase", "Currency", "Purchase Date",
            "Deadline", "Status", "Void Reason",
        ]
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = "A1:N1"

        for column, header in enumerate(headers, 1):
            cell = worksheet.cell(1, column, header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        objects = list(queryset)
        serialized = self.get_serializer(objects, many=True).data
        for row_index, (obj, item) in enumerate(zip(objects, serialized), 2):
            values = [
                item.get("company_name"),
                obj.registration_order.company.national_id,
                item.get("order_number"),
                item.get("payment_instrument_number") or "",
                item.get("purchase_sequence") or "",
                float(item.get("registration_order_amount") or 0),
                float(item.get("amount") or 0),
                float(item.get("order_total_purchased") or 0),
                float(item.get("order_remaining_to_purchase") or 0),
                item.get("currency"),
                item.get("purchase_date"),
                item.get("deadline"),
                "VOID" if item.get("is_void") else "ACTIVE",
                item.get("void_reason") or "",
            ]
            for column, value in enumerate(values, 1):
                worksheet.cell(row_index, column, value)

        for column in range(1, len(headers) + 1):
            width = max(
                len(str(worksheet.cell(row, column).value or ""))
                for row in range(1, worksheet.max_row + 1)
            )
            worksheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 12), 35)

        output = BytesIO()
        workbook.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="currency-purchases.xlsx"'
        response["X-Content-Type-Options"] = "nosniff"
        return response

    @action(detail=True, methods=["post"], url_path="settle-obligation")
    def settle_obligation(self, request, *args, **kwargs):
        purchase = self.get_object()
        if not request.user.has_perm("trade_orders.settle_currencypurchase_obligation"):
            return Response(
                {"detail": "You do not have permission to settle FX obligations."},
                status=status.HTTP_403_FORBIDDEN,
            )
        try:
            updated = mark_obligation_settled(
                purchase=purchase, settled=True, actor=request.user,
                reason=request.data.get("reason", ""),
                reference=request.data.get("reference", ""),
            )
        except DjangoValidationError as exc:
            raise DRFValidationError({"detail": exc.messages})
        return Response(self.get_serializer(updated).data, status=status.HTTP_200_OK)

    @action(detail=True, methods=["post"], url_path="reopen-obligation")
    def reopen_obligation(self, request, *args, **kwargs):
        purchase = self.get_object()
        if not request.user.has_perm("trade_orders.settle_currencypurchase_obligation"):
            return Response(
                {"detail": "You do not have permission to reopen FX obligations."},
                status=status.HTTP_403_FORBIDDEN,
            )
        try:
            updated = mark_obligation_settled(
                purchase=purchase, settled=False, actor=request.user,
                reason=request.data.get("reason", ""),
                reference=request.data.get("reference", ""),
            )
        except DjangoValidationError as exc:
            raise DRFValidationError({"detail": exc.messages})
        return Response(self.get_serializer(updated).data, status=status.HTTP_200_OK)

    def destroy(self, request, *args, **kwargs):
        return Response(
            {
                "detail": (
                    "Currency purchases cannot be deleted. "
                    "Use the correction or void workflow instead."
                )
            },
            status=status.HTTP_405_METHOD_NOT_ALLOWED,
        )

    @action(
        detail=True,
        methods=["post"],
        url_path="void",
    )
    def submit_void(self, request, *args, **kwargs):
        purchase = self.get_object()

        if not request.user.has_perm(
                "trade_orders.void_currencypurchase"
        ):
            return Response(
                {
                    "detail": (
                        "You do not have permission to submit "
                        "a currency purchase void request."
                    )
                },
                status=status.HTTP_403_FORBIDDEN,
            )

        try:
            approval_request = (
                submit_currency_purchase_void(
                    maker=request.user,
                    purchase=purchase,
                    reason=request.data.get(
                        "reason",
                        "",
                    ),
                )
            )
        except DjangoValidationError as exc:
            raise DRFValidationError(
                {"detail": exc.messages}
            )

        return Response(
            {
                "id": str(approval_request.id),
                "status": approval_request.status,
                "operation": approval_request.operation,
                "target_type": approval_request.target_type,
                "target_id": str(
                    approval_request.target_id
                ),
            },
            status=status.HTTP_202_ACCEPTED,
        )

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(
            data=request.data,
        )

        serializer.is_valid(
            raise_exception=True,
        )

        data = serializer.validated_data

        try:
            approval_request = (
                submit_currency_purchase_create(
                    maker=request.user,
                    registration_order=data["registration_order"],
                    amount=data["amount"],
                    currency=data["currency"],
                    purchase_date=data["purchase_date"],
                    remittance_date=data.get("remittance_date"),
                    funding_source_code=data.get("funding_source_code", ""),
                    reason=request.data.get(
                        "reason",
                        "",
                    ),
                )
            )
        except DjangoValidationError as exc:
            raise DRFValidationError(
                {"detail": exc.messages}
            )

        return Response(
            {
                "id": str(approval_request.id),
                "status": approval_request.status,
                "operation": approval_request.operation,
                "target_type": approval_request.target_type,
            },
            status=status.HTTP_202_ACCEPTED,
        )

    def update(self, request, *args, **kwargs):
        return self._update_purchase(
            request=request,
            partial=False,
            *args,
            **kwargs,
        )

    def partial_update(self, request, *args, **kwargs):
        return self._update_purchase(
            request=request,
            partial=True,
            *args,
            **kwargs,
        )

    def _update_purchase(
            self,
            request,
            partial,
            *args,
            **kwargs,
    ):
        purchase = self.get_object()

        serializer = self.get_serializer(
            purchase,
            data=request.data,
            partial=partial,
        )

        serializer.is_valid(
            raise_exception=True,
        )

        data = serializer.validated_data

        requested_currency = data.get(
            "currency",
            purchase.currency,
        )

        if requested_currency != purchase.currency:
            raise DRFValidationError(
                {
                    "currency": (
                        "Currency cannot be changed "
                        "for an existing currency purchase."
                    )
                }
            )

        registration_order = data.get(
            "registration_order",
            purchase.registration_order,
        )

        if (
                registration_order.pk
                != purchase.registration_order_id
        ):
            raise DRFValidationError(
                {
                    "registration_order": (
                        "Registration order cannot be changed "
                        "for an existing currency purchase."
                    )
                }
            )

        try:
            approval_request = (
                submit_currency_purchase_correction(
                    maker=request.user,
                    purchase=purchase,
                    amount=data.get(
                        "amount",
                        purchase.amount,
                    ),
                    purchase_date=data.get(
                        "purchase_date",
                        purchase.purchase_date,
                    ),
                    remittance_date=data.get(
                        "remittance_date",
                        purchase.remittance_date,
                    ),
                    funding_source_code=data.get(
                        "funding_source_code",
                        purchase.funding_source_code,
                    ),
                    reason=request.data.get(
                        "reason",
                        "",
                    ),
                )
            )
        except DjangoValidationError as exc:
            raise DRFValidationError(
                {"detail": exc.messages}
            )

        return Response(
            {
                "id": str(approval_request.id),
                "status": approval_request.status,
                "operation": approval_request.operation,
                "target_type": approval_request.target_type,
                "target_id": str(
                    approval_request.target_id
                ),
            },
            status=status.HTTP_202_ACCEPTED,
        )

class ShipmentPartViewSet(viewsets.ModelViewSet):
    permission_classes = (
        TradeFlowModelPermissions,
    )

    queryset = (
        ShipmentPart.objects
        .select_related(
            "currency_purchase",
            "currency_purchase__registration_order",
            "currency_purchase__registration_order__company",
        )
        .all()
    )

    serializer_class = ShipmentPartSerializer

    search_fields = (
        "reference_number",
        "currency_purchase__registration_order__order_number",
        "currency_purchase__registration_order__company__name",
        "currency_purchase__registration_order__company__national_id",
    )

    ordering_fields = (
        "shipment_date",
        "received_date",
        "amount",
        "created_at",
    )

    @action(detail=True, methods=["post"], url_path="settle-obligation")
    def settle_obligation(self, request, *args, **kwargs):
        purchase = self.get_object()
        if not request.user.has_perm("trade_orders.settle_currencypurchase_obligation"):
            return Response(
                {"detail": "You do not have permission to settle FX obligations."},
                status=status.HTTP_403_FORBIDDEN,
            )
        try:
            updated = mark_obligation_settled(
                purchase=purchase, settled=True, actor=request.user,
                reason=request.data.get("reason", ""),
                reference=request.data.get("reference", ""),
            )
        except DjangoValidationError as exc:
            raise DRFValidationError({"detail": exc.messages})
        return Response(self.get_serializer(updated).data, status=status.HTTP_200_OK)

    @action(detail=True, methods=["post"], url_path="reopen-obligation")
    def reopen_obligation(self, request, *args, **kwargs):
        purchase = self.get_object()
        if not request.user.has_perm("trade_orders.settle_currencypurchase_obligation"):
            return Response(
                {"detail": "You do not have permission to reopen FX obligations."},
                status=status.HTTP_403_FORBIDDEN,
            )
        try:
            updated = mark_obligation_settled(
                purchase=purchase, settled=False, actor=request.user,
                reason=request.data.get("reason", ""),
                reference=request.data.get("reference", ""),
            )
        except DjangoValidationError as exc:
            raise DRFValidationError({"detail": exc.messages})
        return Response(self.get_serializer(updated).data, status=status.HTTP_200_OK)

    def destroy(self, request, *args, **kwargs):
        return Response(
            {
                "detail": (
                    "Shipment parts cannot be deleted directly. "
                    "Use the correction or void workflow instead."
                )
            },
            status=status.HTTP_405_METHOD_NOT_ALLOWED,
        )

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(
            data=request.data,
        )

        serializer.is_valid(
            raise_exception=True,
        )

        data = serializer.validated_data

        try:
            approval_request = (
                submit_shipment_part_create(
                    maker=request.user,
                    currency_purchase=data[
                        "currency_purchase"
                    ],
                    amount=data["amount"],
                    shipment_date=data.get(
                        "shipment_date"
                    ),
                    received_date=data.get(
                        "received_date"
                    ),
                    reference_number=data.get(
                        "reference_number",
                        "",
                    ),
                    notes=data.get(
                        "notes",
                        "",
                    ),
                    reason=request.data.get(
                        "reason",
                        "",
                    ),
                )
            )
        except DjangoValidationError as exc:
            raise DRFValidationError(
                {"detail": exc.messages}
            )

        return Response(
            {
                "id": str(approval_request.id),
                "status": approval_request.status,
                "operation": approval_request.operation,
                "target_type":
                    approval_request.target_type,
            },
            status=status.HTTP_202_ACCEPTED,
        )

    def update(self, request, *args, **kwargs):
        return self._update_shipment(
            request=request,
            partial=False,
            *args,
            **kwargs,
        )

    def partial_update(
            self,
            request,
            *args,
            **kwargs,
    ):
        return self._update_shipment(
            request=request,
            partial=True,
            *args,
            **kwargs,
        )

    def _update_shipment(
            self,
            request,
            partial,
            *args,
            **kwargs,
    ):
        shipment = self.get_object()

        if (
                "currency_purchase" in request.data
                and str(request.data["currency_purchase"])
                != str(shipment.currency_purchase_id)
        ):
            raise DRFValidationError(
                {
                    "currency_purchase": (
                        "Currency purchase cannot be changed "
                        "for an existing shipment part."
                    )
                }
            )

        serializer = self.get_serializer(
            shipment,
            data=request.data,
            partial=partial,
        )

        serializer.is_valid(
            raise_exception=True,
        )

        data = serializer.validated_data

        try:
            approval_request = (
                submit_shipment_part_correction(
                    maker=request.user,
                    shipment=shipment,
                    amount=data.get(
                        "amount",
                        shipment.amount,
                    ),
                    shipment_date=data.get(
                        "shipment_date",
                        shipment.shipment_date,
                    ),
                    received_date=data.get(
                        "received_date",
                        shipment.received_date,
                    ),
                    reference_number=data.get(
                        "reference_number",
                        shipment.reference_number,
                    ),
                    notes=data.get(
                        "notes",
                        shipment.notes,
                    ),
                    reason=request.data.get(
                        "reason",
                        "",
                    ),
                )
            )
        except DjangoValidationError as exc:
            raise DRFValidationError(
                {"detail": exc.messages}
            )

        return Response(
            {
                "id": str(approval_request.id),
                "status": approval_request.status,
                "operation": approval_request.operation,
                "target_type":
                    approval_request.target_type,
                "target_id": str(
                    approval_request.target_id
                ),
            },
            status=status.HTTP_202_ACCEPTED,
        )

    @action(
        detail=True,
        methods=["post"],
        url_path="void",
    )
    def submit_void(
            self,
            request,
            *args,
            **kwargs,
    ):
        shipment = self.get_object()

        try:
            approval_request = (
                submit_shipment_part_void(
                    maker=request.user,
                    shipment=shipment,
                    reason=request.data.get(
                        "reason",
                        "",
                    ),
                )
            )
        except DjangoValidationError as exc:
            raise DRFValidationError(
                {"detail": exc.messages}
            )

        return Response(
            {
                "id": str(approval_request.id),
                "status": approval_request.status,
                "operation": approval_request.operation,
                "target_type":
                    approval_request.target_type,
                "target_id": str(
                    approval_request.target_id
                ),
            },
            status=status.HTTP_202_ACCEPTED,
        )

class DashboardSummaryAPIView(APIView):
    permission_classes = (
        CanViewTradeDashboard,
    )

    def get(self, request):
        summary = get_dashboard_summary()

        return Response(
            summary,
            status=status.HTTP_200_OK,
        )

class RegulatoryRuleViewSet(viewsets.ModelViewSet):
    permission_classes = (TradeFlowModelPermissions,)
    queryset = RegulatoryRule.objects.all()
    serializer_class = RegulatoryRuleSerializer
    filterset_fields = ("operation_type", "activity_type", "goods_category_code", "funding_source_code", "deadline_kind", "is_active")
    ordering_fields = ("priority", "effective_from", "code")


class RegulatoryDeadlineViewSet(viewsets.ReadOnlyModelViewSet):
    permission_classes = (TradeFlowModelPermissions,)
    queryset = RegulatoryDeadline.objects.select_related(
        "currency_purchase", "applied_rule"
    ).all()
    serializer_class = RegulatoryDeadlineSerializer
    filterset_fields = ("currency_purchase", "deadline_kind")
    ordering_fields = ("effective_deadline", "original_deadline", "deadline_kind")


class CustomsClearanceViewSet(viewsets.ModelViewSet):
    permission_classes = (TradeFlowModelPermissions,)
    queryset = CustomsClearance.objects.select_related("currency_purchase").all()
    serializer_class = CustomsClearanceSerializer
    filterset_fields = ("currency_purchase", "status")
    search_fields = ("declaration_number",)

    def perform_create(self, serializer):
        obj = serializer.save()
        refresh_obligation_status(obj.currency_purchase)

    def perform_update(self, serializer):
        obj = serializer.save()
        refresh_obligation_status(obj.currency_purchase)


class DeadlineExtensionViewSet(viewsets.ReadOnlyModelViewSet):
    permission_classes = (TradeFlowModelPermissions,)
    queryset = DeadlineExtension.objects.select_related("currency_purchase", "approved_by").all()
    serializer_class = DeadlineExtensionSerializer
    filterset_fields = ("currency_purchase",)

    @action(detail=False, methods=["post"], url_path="apply")
    def apply_extension(self, request):
        purchase = get_object_or_404(CurrencyPurchase, pk=request.data.get("currency_purchase"))
        if not request.user.has_perm("trade_orders.extend_currencypurchase_deadline"):
            return Response({"detail": "You do not have permission to extend deadlines."}, status=status.HTTP_403_FORBIDDEN)
        serializer = DeadlineExtensionSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        event = extend_deadline(
            purchase=purchase,
            new_deadline=serializer.validated_data["new_deadline"],
            reason=serializer.validated_data["reason"],
            reference=serializer.validated_data.get("reference", ""),
            approved_by=request.user,
            deadline_kind=request.data.get("deadline_kind", RegulatoryRule.DeadlineKind.IMPORT_CLEARANCE),
        )
        return Response(self.get_serializer(event).data, status=status.HTTP_201_CREATED)
