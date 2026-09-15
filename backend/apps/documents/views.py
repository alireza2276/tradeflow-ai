from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from rest_framework.decorators import action

from django.core.exceptions import ValidationError as DjangoValidationError
from rest_framework import status, viewsets
from rest_framework.exceptions import ValidationError as DRFValidationError
from rest_framework.response import Response

from apps.authentication.permissions import (
    TradeFlowModelPermissions,
)
from apps.documents.models import Invoice
from apps.documents.serializers import InvoiceSerializer
from apps.documents.services.invoice_service import (
    create_invoice,
    update_invoice,
)


class InvoiceViewSet(viewsets.ModelViewSet):
    permission_classes = (
        TradeFlowModelPermissions,
    )

    queryset = (
        Invoice.objects
        .select_related(
            "shipment_part",
            "shipment_part__currency_purchase",
            "shipment_part__currency_purchase__registration_order",
            "shipment_part__currency_purchase__registration_order__company",
            "shipment_part__currency_purchase__registration_order__payment_instrument",
        )
        .all()
    )

    serializer_class = InvoiceSerializer

    search_fields = (
        "shipment_part__reference_number",
        "shipment_part__currency_purchase__registration_order__order_number",
        "shipment_part__currency_purchase__registration_order__company__name",
        "shipment_part__currency_purchase__registration_order__company__national_id",
        "shipment_part__currency_purchase__registration_order__payment_instrument__instrument_number",
    )

    ordering_fields = (
        "submission_date",
        "total_amount",
        "created_at",
    )

    def get_queryset(self):
        queryset = super().get_queryset()
        params = self.request.query_params
        text_filters = (
            ("company", "shipment_part__currency_purchase__registration_order__company__name__icontains"),
            ("national_id", "shipment_part__currency_purchase__registration_order__company__national_id__icontains"),
            ("order_number", "shipment_part__currency_purchase__registration_order__order_number__icontains"),
            ("instrument_number", "shipment_part__currency_purchase__registration_order__payment_instrument__instrument_number__icontains"),
            ("reference_number", "shipment_part__reference_number__icontains"),
        )
        for param, lookup in text_filters:
            value = params.get(param, "").strip()
            if value:
                queryset = queryset.filter(**{lookup: value})

        for param, lookup in (
            ("purchase_date_from", "shipment_part__currency_purchase__purchase_date__gte"),
            ("purchase_date_to", "shipment_part__currency_purchase__purchase_date__lte"),
            ("deadline_from", "shipment_part__currency_purchase__deadline__gte"),
            ("deadline_to", "shipment_part__currency_purchase__deadline__lte"),
            ("submission_date_from", "submission_date__gte"),
            ("submission_date_to", "submission_date__lte"),
        ):
            value = params.get(param, "").strip()
            if value:
                queryset = queryset.filter(**{lookup: value})
        return queryset

    @action(detail=False, methods=["get"], url_path="export-xlsx")
    def export_xlsx(self, request):
        objects = list(self.filter_queryset(self.get_queryset()))
        data = self.get_serializer(objects, many=True).data
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Invoices"
        headers = [
            "Company", "Registration Order", "Payment Instrument",
            "Document Part", "Shipment Reference", "Purchase Date",
            "Deadline", "Purchase Tranche Amount", "Total Purchased",
            "FOB", "Freight", "Invoice Total", "Remaining Documents",
            "Currency", "Submission Date",
        ]
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = "A1:O1"
        for column, header in enumerate(headers, 1):
            cell = worksheet.cell(1, column, header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
        for row_index, item in enumerate(data, 2):
            values = [
                item.get("company_name"), item.get("order_number"),
                item.get("payment_instrument_number") or "",
                item.get("document_part_number"),
                item.get("shipment_reference_number") or "",
                item.get("purchase_date"), item.get("deadline"),
                float(item.get("currency_purchase_amount") or 0),
                float(item.get("order_total_purchased") or 0),
                float(item.get("fob_amount") or 0),
                float(item.get("freight_amount") or 0),
                float(item.get("total_amount") or 0),
                float(item.get("remaining_amount") or 0),
                item.get("order_currency"), item.get("submission_date"),
            ]
            for column, value in enumerate(values, 1):
                worksheet.cell(row_index, column, value)
        for column in range(1, len(headers) + 1):
            width = max(len(str(worksheet.cell(row, column).value or "")) for row in range(1, worksheet.max_row + 1))
            worksheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 12), 35)
        output = BytesIO()
        workbook.save(output)
        response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = 'attachment; filename="invoices.xlsx"'
        response["X-Content-Type-Options"] = "nosniff"
        return response

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(
            data=request.data,
        )

        serializer.is_valid(
            raise_exception=True,
        )

        data = serializer.validated_data

        try:
            invoice = create_invoice(
                shipment_part=data["shipment_part"],
                fob_amount=data["fob_amount"],
                freight_amount=data["freight_amount"],
                submission_date=data["submission_date"],
            )
        except DjangoValidationError as exc:
            raise DRFValidationError(
                {"detail": exc.messages}
            )

        output_serializer = self.get_serializer(
            invoice,
        )

        return Response(
            output_serializer.data,
            status=status.HTTP_201_CREATED,
        )

    def update(self, request, *args, **kwargs):
        partial = kwargs.pop(
            "partial",
            False,
        )

        invoice = self.get_object()

        serializer = self.get_serializer(
            invoice,
            data=request.data,
            partial=partial,
        )

        serializer.is_valid(
            raise_exception=True,
        )

        data = serializer.validated_data

        shipment_part = data.get(
            "shipment_part",
            invoice.shipment_part,
        )

        if shipment_part.pk != invoice.shipment_part_id:
            raise DRFValidationError(
                {
                    "shipment_part": (
                        "Shipment part cannot be changed "
                        "after invoice creation."
                    )
                }
            )

        try:
            updated_invoice = update_invoice(
                invoice=invoice,
                fob_amount=data.get(
                    "fob_amount",
                ),
                freight_amount=data.get(
                    "freight_amount",
                ),
                submission_date=data.get(
                    "submission_date",
                ),
            )

        except DjangoValidationError as exc:
            raise DRFValidationError(
                {"detail": exc.messages}
            )

        output_serializer = self.get_serializer(
            updated_invoice,
        )

        return Response(
            output_serializer.data
        )

    def destroy(self, request, *args, **kwargs):
        return Response(
            {
                "detail": (
                    "Invoice deletion is not allowed. "
                    "Use a correction or void workflow instead."
                )
            },
            status=status.HTTP_405_METHOD_NOT_ALLOWED,
        )
